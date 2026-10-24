"""Convert TPW's "Collective Impact Report" workbook into the seed CSVs.

    python3 -I scripts/convert_tpw_xlsx.py "<path to TPW - Collective Impact Report.xlsx>"

Writes data/real/facilities.csv and data/real/program_runs.csv. The workbook
itself is not committed. Rows marked "SAMPLE DATA" in Notes and blank template
rows are skipped; "TBD" and blank numbers are kept as blanks (the dashboard
lists them as incomplete rather than guessing).

Facility flags come from published sources, not from the workbook:
- right4_underserved: red-font "underserved institutions" in Attachment A of
  CDCR's RIGHT Grant 4.0 RFA (released 2026-10-01, due 2026-11-30) — worth
  +0.5 qualifying points.
- womens: CDCR's women's institutions (CCWF, CIW).
"""
import csv
import datetime as dt
import os
import re
import sys

import openpyxl

RIGHT4_UNDERSERVED = {"ASP", "CCWF", "CIW", "CTF", "KVSP", "NKSP", "SAC"}
WOMENS = {"CCWF", "CIW"}
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "real")


def code_of(name):
    m = re.search(r"\(([A-Z]+)", name or "")
    return m.group(1) if m else None


def num(v):
    return "" if not isinstance(v, (int, float)) else (str(int(v)) if float(v).is_integer() else str(v))


def day(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) else ""


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    facilities = [r[0] for r in wb["Constants"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    with open(os.path.join(OUT, "facilities.csv"), "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["code", "name", "right4_underserved", "womens"])
        for name in facilities:
            c = code_of(name)
            w.writerow([c, name.strip(), int(c in RIGHT4_UNDERSERVED), int(c in WOMENS)])

    rows = list(wb["Data Entry Sheet"].iter_rows(values_only=True))
    kept = skipped = 0
    with open(os.path.join(OUT, "program_runs.csv"), "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["source_row", "facility_code", "yard", "organization", "program", "status",
                    "start_date", "end_date", "total_hours", "participants", "completed", "notes"])
        for i, r in enumerate(rows[1:], start=2):
            if not r or not r[2]:
                continue
            if (r[12] or "").strip().upper() == "SAMPLE DATA":
                skipped += 1
                continue
            w.writerow([i, code_of(r[0]) or "", (r[1] or "").strip(), r[2].strip(), (r[3] or "").strip(),
                        (r[4] or "").strip(), day(r[5]), day(r[6]), num(r[7]), num(r[8]), num(r[9]),
                        (r[12] or "").strip()])
            kept += 1
    print(f"facilities: {len(facilities)}; program runs kept: {kept}; sample rows skipped: {skipped}")


if __name__ == "__main__":
    main(sys.argv[1])
